from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from express_app.core.account_balance import (  # noqa: E402
    CustomerBalanceRecord,
    collect_account_balance_dashboard,
)
from express_app.core.calculator import (  # noqa: E402
    CUSTOMER_ABNORMAL_DEDUCTION_HEADERS,
    CUSTOMER_ABNORMAL_DEDUCTION_SHEET,
    CUSTOMER_HISTORY_SHEET,
    CUSTOMER_PAYMENT_HEADERS,
    CUSTOMER_PAYMENT_SHEET,
    customer_history_summary_file_name,
)


class AccountBalanceDashboardTest(unittest.TestCase):
    def _write_history_workbook(
        self,
        customer_dir: Path,
        customer: str,
        rows: list[tuple[date, float]],
        payments: list[tuple[date, float]] | None = None,
        abnormal_deductions: list[tuple[date, float, str]] | None = None,
    ) -> None:
        workbook = openpyxl.Workbook()
        history_sheet = workbook.active
        history_sheet.title = CUSTOMER_HISTORY_SHEET
        history_sheet.append(
            [
                "日期",
                "单数",
                "总重量",
                "今日快递总消费",
                "累计快递费用",
                "今日收款",
                "累计收款",
                "当前余额",
            ]
        )
        cumulative = 0.0
        for row_index, (shipping_date, fee) in enumerate(rows, start=2):
            cumulative += fee
            history_sheet.append([shipping_date, 1, 1.0, fee, cumulative, "", "", ""])

        payment_sheet = workbook.create_sheet(CUSTOMER_PAYMENT_SHEET)
        payment_sheet.append(CUSTOMER_PAYMENT_HEADERS)
        for paid_date, amount in payments or []:
            payment_sheet.append([paid_date, "微信", "小李", amount, paid_date])

        if abnormal_deductions is not None:
            deduction_sheet = workbook.create_sheet(CUSTOMER_ABNORMAL_DEDUCTION_SHEET)
            deduction_sheet.append(CUSTOMER_ABNORMAL_DEDUCTION_HEADERS)
            for deduction_date, amount, remark in abnormal_deductions:
                deduction_sheet.append([deduction_date, amount, remark])

        workbook.save(customer_dir / customer_history_summary_file_name(customer))

    def test_collects_customer_balance_from_history_and_payment_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            split_dir = Path(temp_dir)
            customer_a = split_dir / "客户A"
            customer_b = split_dir / "客户B"
            customer_a.mkdir()
            customer_b.mkdir()
            self._write_history_workbook(
                customer_a,
                "客户A",
                [(date(2026, 4, 1), 120.5), (date(2026, 4, 2), 80.0)],
                [(date(2026, 4, 1), 500.0)],
            )
            self._write_history_workbook(
                customer_b,
                "客户B",
                [(date(2026, 4, 3), 300.0)],
                [(date(2026, 4, 3), 100.0)],
            )

            dashboard = collect_account_balance_dashboard(split_dir)

            self.assertEqual(dashboard.customer_count, 2)
            self.assertAlmostEqual(dashboard.total_consumed, 500.5)
            self.assertAlmostEqual(dashboard.total_paid, 600.0)
            self.assertAlmostEqual(dashboard.total_balance, 99.5)
            self.assertEqual(dashboard.debtor_count, 1)
            self.assertEqual([record.customer for record in dashboard.records], ["客户B", "客户A"])

            debtor = dashboard.records[0]
            self.assertIsInstance(debtor, CustomerBalanceRecord)
            self.assertEqual(debtor.customer, "客户B")
            self.assertEqual(debtor.last_date, date(2026, 4, 3))
            self.assertAlmostEqual(debtor.total_consumed, 300.0)
            self.assertAlmostEqual(debtor.total_paid, 100.0)
            self.assertAlmostEqual(debtor.current_balance, -200.0)
            self.assertEqual(debtor.status, "欠款")
            self.assertEqual(debtor.history_file, customer_b / customer_history_summary_file_name("客户B"))
            self.assertEqual(debtor.customer_dir, customer_b)

    def test_account_balance_deducts_abnormal_deduction_records(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            split_dir = Path(temp_dir)
            customer_dir = split_dir / "客户A"
            customer_dir.mkdir()
            self._write_history_workbook(
                customer_dir,
                "客户A",
                [(date(2026, 4, 2), 200.0)],
                [(date(2026, 4, 2), 500.0)],
                [(date(2026, 4, 2), 80.0, "破损扣款")],
            )

            dashboard = collect_account_balance_dashboard(split_dir)

            self.assertEqual(dashboard.customer_count, 1)
            record = dashboard.records[0]
            self.assertAlmostEqual(record.total_consumed, 200.0)
            self.assertAlmostEqual(record.total_paid, 500.0)
            self.assertAlmostEqual(record.current_balance, 220.0)
            self.assertAlmostEqual(dashboard.total_balance, 220.0)

    def test_missing_split_directory_returns_error_dashboard(self) -> None:
        missing_dir = Path("/tmp/aisong-missing-balance-dir")

        dashboard = collect_account_balance_dashboard(missing_dir)

        self.assertEqual(dashboard.customer_count, 0)
        self.assertEqual(dashboard.records, [])
        self.assertIn("客户明细目录不存在", dashboard.errors[0])

    def test_skips_invalid_history_workbook_and_reports_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            split_dir = Path(temp_dir)
            customer_dir = split_dir / "客户A"
            customer_dir.mkdir()
            workbook = openpyxl.Workbook()
            workbook.save(customer_dir / customer_history_summary_file_name("客户A"))

            dashboard = collect_account_balance_dashboard(split_dir)

            self.assertEqual(dashboard.records, [])
            self.assertEqual(dashboard.customer_count, 0)
            self.assertIn("客户A", dashboard.errors[0])
            self.assertIn("缺少 sheet", dashboard.errors[0])


if __name__ == "__main__":
    unittest.main()
