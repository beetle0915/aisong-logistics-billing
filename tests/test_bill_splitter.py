from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from express_app.core.bill_splitter import (  # noqa: E402
    scan_bill_split_directory,
    split_bills_by_field,
)


def write_workbook(path: Path, sheets: dict[str, tuple[list[str], list[list[object]]]]) -> None:
    workbook = openpyxl.Workbook()
    for index, (sheet_name, (headers, rows)) in enumerate(sheets.items()):
        ws = workbook.active if index == 0 else workbook.create_sheet()
        ws.title = sheet_name
        ws.append(headers)
        for row in rows:
            ws.append(row)
    workbook.save(path)


def read_headers(ws: openpyxl.worksheet.worksheet.Worksheet) -> list[object]:
    return [cell.value for cell in ws[1]]


class BillSplitterTest(unittest.TestCase):
    def test_scan_bill_split_directory_is_read_only_and_reports_direct_bill_detail_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            nested_dir = input_dir / "nested"
            nested_dir.mkdir()
            output_dir = input_dir.parent / "账单拆分结果【拆分字段：经手人】"
            headers = ["运单号码", "经手人", "费用", "渠道"]

            write_workbook(input_dir / "A.xlsx", {"账单明细": (headers, [["YD001", "张三", 12.5, "顺丰"]])})
            write_workbook(
                input_dir / "B.xlsx",
                {"账单明细": (["渠道", "费用", "经手人", "运单号码"], [["顺丰", 20, "李四", "YD002"]])},
            )
            write_workbook(input_dir / "缺明细.xlsx", {"Sheet2": (headers, [["YD003", "王五", 3, "顺丰"]])})
            write_workbook(input_dir / ".隐藏.xlsx", {"账单明细": (headers, [["YD004", "隐藏", 4, "顺丰"]])})
            write_workbook(input_dir / "~$临时.xlsx", {"账单明细": (headers, [["YD005", "临时", 5, "顺丰"]])})
            write_workbook(nested_dir / "子目录.xlsx", {"账单明细": (headers, [["YD006", "子目录", 6, "顺丰"]])})

            result = scan_bill_split_directory(input_dir)

            self.assertEqual(result.output_dir, output_dir)
            self.assertFalse(output_dir.exists())
            self.assertEqual(result.common_headers, ["运单号码", "经手人", "费用", "渠道"])
            self.assertEqual([file.path.name for file in result.files], ["A.xlsx", "B.xlsx", "缺明细.xlsx"])
            self.assertEqual([file.sheet_name for file in result.files], ["账单明细", "账单明细", "账单明细"])
            self.assertEqual([file.total_rows for file in result.files], [1, 1, 0])
            self.assertEqual(result.files[-1].error, "缺少账单明细")
            self.assertEqual(result.errors, [])
            self.assertTrue(any("缺明细.xlsx" in log and "缺少账单明细" in log for log in result.logs))

    def test_scan_bill_split_directory_reports_missing_common_fields_without_writing(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            output_dir = input_dir.parent / "账单拆分结果【拆分字段：经手人】"
            write_workbook(
                input_dir / "A.xlsx",
                {"账单明细": (["运单号码", "经手人", "费用"], [["YD001", "张三", 12.5]])},
            )
            write_workbook(
                input_dir / "B.xlsx",
                {"账单明细": (["运单号码", "费用"], [["YD002", 20]])},
            )

            result = scan_bill_split_directory(input_dir)

            self.assertFalse(output_dir.exists())
            self.assertEqual(result.common_headers, ["运单号码", "费用"])
            self.assertEqual([file.path.name for file in result.files], ["A.xlsx", "B.xlsx"])
            self.assertIn("共同表头缺少拆分字段：经手人", result.errors)

    def test_bad_xlsx_file_is_reported_without_blocking_valid_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            headers = ["运单号码", "经手人", "费用"]
            write_workbook(input_dir / "A.xlsx", {"账单明细": (headers, [["YD001", "张三", 12.5]])})
            (input_dir / "坏文件.xlsx").write_text("not a real workbook", encoding="utf-8")

            scan = scan_bill_split_directory(input_dir)
            result = split_bills_by_field(input_dir)

            self.assertEqual(scan.common_headers, headers)
            self.assertEqual([file.path.name for file in scan.files], ["A.xlsx", "坏文件.xlsx"])
            failed_files = [file for file in scan.files if file.error]
            self.assertEqual([file.path.name for file in failed_files], ["坏文件.xlsx"])
            self.assertTrue(any("坏文件.xlsx" in log and "读取失败" in log for log in scan.logs))
            self.assertEqual([path.name for path in result.output_paths], ["经手人_张三.xlsx"])

    def test_scan_ignores_style_only_tail_rows_that_pollute_excel_max_row(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            workbook = openpyxl.Workbook()
            ws = workbook.active
            ws.title = "账单明细"
            ws.append(["运单号码", "经手人", "费用"])
            ws.append(["YD001", "张三", 12.5])
            ws.cell(row=5000, column=1).fill = PatternFill("solid", fgColor="FFFFFF")
            workbook.save(input_dir / "明细.xlsx")

            result = scan_bill_split_directory(input_dir)

            self.assertEqual([file.total_rows for file in result.files], [1])

    def test_scans_only_direct_visible_xlsx_files_and_exact_bill_detail_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            nested_dir = input_dir / "nested"
            nested_dir.mkdir()
            headers = ["运单号码", "经手人", "费用"]

            write_workbook(
                input_dir / "正确.xlsx",
                {
                    "Sheet2": (headers, [["IGNORED", "Sheet2人员", 1]]),
                    "账单明细": (headers, [["YD001", "张三", 12.5]]),
                    "账单总览": (headers, [["IGNORED2", "总览人员", 2]]),
                },
            )
            write_workbook(input_dir / "缺明细.xlsx", {"Sheet2": (headers, [["YD002", "李四", 3]])})
            write_workbook(input_dir / ".隐藏.xlsx", {"账单明细": (headers, [["YD003", "隐藏人员", 4]])})
            write_workbook(input_dir / "~$临时.xlsx", {"账单明细": (headers, [["YD004", "临时人员", 5]])})
            write_workbook(input_dir / "非xlsx.xlsm", {"账单明细": (headers, [["YD005", "宏人员", 6]])})
            write_workbook(nested_dir / "子目录.xlsx", {"账单明细": (headers, [["YD006", "子目录人员", 7]])})
            (input_dir / "文本.txt").write_text("not excel", encoding="utf-8")

            result = split_bills_by_field(input_dir)

            self.assertEqual(result.common_headers, headers)
            self.assertEqual([path.name for path in result.output_paths], ["经手人_张三.xlsx"])
            workbook = openpyxl.load_workbook(result.output_paths[0], data_only=True)
            self.assertEqual(workbook.sheetnames, ["正确"])
            ws = workbook["正确"]
            self.assertEqual(ws["A2"].value, "YD001")
            self.assertEqual(ws.max_row, 2)

    def test_common_headers_keep_first_valid_file_order_and_include_split_field(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            write_workbook(
                input_dir / "A.xlsx",
                {
                    "账单明细": (
                        ["运单号码", "经手人", "费用", "渠道", "只在A"],
                        [["YD001", "张三", 12.5, "顺丰", "A"]],
                    )
                },
            )
            write_workbook(
                input_dir / "B.xlsx",
                {
                    "账单明细": (
                        ["渠道", "费用", "经手人", "运单号码", "只在B"],
                        [["顺丰", 20, "李四", "YD002", "B"]],
                    )
                },
            )

            result = split_bills_by_field(input_dir)

            self.assertEqual(result.common_headers, ["运单号码", "经手人", "费用", "渠道"])

    def test_split_value_across_two_source_files_writes_one_workbook_with_two_sheets(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "五月账单"
            input_dir.mkdir()
            headers = ["运单号码", "经手人", "费用"]
            write_workbook(input_dir / "第一份.xlsx", {"账单明细": (headers, [["YD001", "张三", 12.5]])})
            write_workbook(input_dir / "第二份.xlsx", {"账单明细": (headers, [["YD002", "张三", 20]])})

            result = split_bills_by_field(input_dir)

            self.assertEqual(result.output_dir, input_dir.parent / "五月账单拆分结果【拆分字段：经手人】")
            self.assertEqual([path.name for path in result.output_paths], ["经手人_张三.xlsx"])
            workbook = openpyxl.load_workbook(result.output_paths[0], data_only=True)
            self.assertEqual(workbook.sheetnames, ["第一份", "第二份"])
            self.assertEqual(workbook["第一份"]["A2"].value, "YD001")
            self.assertEqual(workbook["第二份"]["A2"].value, "YD002")
            self.assertEqual(workbook["第一份"]["D1"].value, "来源文件")
            self.assertEqual(workbook["第一份"]["E1"].value, "来源行号")

    def test_split_bills_by_selected_field_instead_of_default_handler_field(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            headers = ["运单号码", "省份", "经手人", "费用"]
            write_workbook(
                input_dir / "明细.xlsx",
                {
                    "账单明细": (
                        headers,
                        [
                            ["YD001", "浙江", "张三", 12.5],
                            ["YD002", "江苏", "张三", 20],
                        ],
                    )
                },
            )

            result = split_bills_by_field(input_dir, "省份")

            self.assertEqual(result.output_dir, input_dir.parent / "账单拆分结果【拆分字段：省份】")
            self.assertEqual([path.name for path in result.output_paths], ["省份_江苏.xlsx", "省份_浙江.xlsx"])
            self.assertEqual([output.split_value for output in result.outputs], ["江苏", "浙江"])

    def test_summary_rows_and_blank_or_invalid_split_values_do_not_generate_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            headers = ["运单号码", "经手人", "费用"]
            write_workbook(
                input_dir / "明细.xlsx",
                {
                    "账单明细": (
                        headers,
                        [
                            ["YD001", "张三", 12.5],
                            ["合计", "159985.00", 159985.00],
                            ["小计", "23885.32", 23885.32],
                            ["YD002", "", 3],
                            ["YD003", "#N/A", 4],
                            ["YD004", "N/A", 5],
                            ["YD005", "无", 6],
                            [None, "王五", 6],
                            [None, None, None],
                        ],
                    )
                },
            )

            result = split_bills_by_field(input_dir)

            self.assertEqual([path.name for path in result.output_paths], ["经手人_张三.xlsx"])
            output_names = "\n".join(path.name for path in result.output_paths)
            self.assertNotIn("159985.00", output_names)
            self.assertNotIn("23885.32", output_names)

    def test_existing_output_directory_overwrites_matching_file_without_clearing_directory(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            output_dir = input_dir.parent / "账单拆分结果【拆分字段：经手人】"
            output_dir.mkdir()
            keep_file = output_dir / "保留.txt"
            keep_file.write_text("keep", encoding="utf-8")
            stale_file = output_dir / "经手人_张三.xlsx"
            write_workbook(stale_file, {"旧内容": (["A"], [["stale"]])})

            headers = ["运单号码", "经手人", "费用"]
            write_workbook(input_dir / "明细.xlsx", {"账单明细": (headers, [["YD001", "张三", 12.5]])})

            result = split_bills_by_field(input_dir)

            self.assertTrue(keep_file.exists())
            self.assertEqual(result.output_paths, [stale_file])
            workbook = openpyxl.load_workbook(stale_file, data_only=True)
            self.assertEqual(workbook.sheetnames, ["明细"])
            self.assertEqual(workbook["明细"]["A2"].value, "YD001")
            self.assertNotIn("旧内容", workbook.sheetnames)

    def test_file_and_sheet_names_are_sanitized_and_sheet_conflicts_get_suffixes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            headers = ["运单号码", "经手人", "费用"]
            split_value = 'A/B:C*D?E"F<G>H|I'
            write_workbook(input_dir / "客户:一.xlsx", {"账单明细": (headers, [["YD001", split_value, 12.5]])})
            write_workbook(input_dir / "客户?一.xlsx", {"账单明细": (headers, [["YD002", split_value, 20]])})

            result = split_bills_by_field(input_dir)

            self.assertEqual([path.name for path in result.output_paths], ["经手人_A_B_C_D_E_F_G_H_I.xlsx"])
            workbook = openpyxl.load_workbook(result.output_paths[0], data_only=True)
            self.assertEqual(workbook.sheetnames, ["客户_一", "客户_一_2"])

    def test_sanitized_file_name_collisions_get_suffixes(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            headers = ["运单号码", "经手人", "费用"]
            write_workbook(
                input_dir / "明细.xlsx",
                {
                    "账单明细": (
                        headers,
                        [
                            ["YD001", "A/B", 12.5],
                            ["YD002", "A:B", 20],
                        ],
                    )
                },
            )

            result = split_bills_by_field(input_dir)

            self.assertEqual([path.name for path in result.output_paths], ["经手人_A_B.xlsx", "经手人_A_B_2.xlsx"])
            self.assertEqual([output.split_value for output in result.outputs], ["A/B", "A:B"])
            self.assertEqual([output.row_count for output in result.outputs], [1, 1])

    def test_split_result_outputs_keep_real_split_value_and_row_count(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            input_dir = Path(temp_dir_text) / "账单"
            input_dir.mkdir()
            headers = ["运单号码", "经手人", "费用"]
            write_workbook(
                input_dir / "明细.xlsx",
                {
                    "账单明细": (
                        headers,
                        [
                            ["YD001", "A/B", 12.5],
                            ["YD002", "A/B", 20],
                            ["YD003", "超长" * 80, 30],
                        ],
                    )
                },
            )

            result = split_bills_by_field(input_dir)

            self.assertEqual(len(result.outputs), 2)
            self.assertEqual(result.output_paths, [output.output_path for output in result.outputs])
            self.assertEqual(result.outputs[0].split_value, "A/B")
            self.assertEqual(result.outputs[0].row_count, 2)
            self.assertEqual(result.outputs[1].split_value, "超长" * 80)
            self.assertEqual(result.outputs[1].row_count, 1)
            self.assertNotEqual(result.outputs[1].split_value, result.outputs[1].output_path.stem.removeprefix("经手人_"))


if __name__ == "__main__":
    unittest.main()
