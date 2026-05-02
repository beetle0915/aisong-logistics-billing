from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from express_app.core import calculator  # noqa: E402
from express_app.core.calculator import (  # noqa: E402
    load_price_template_workbook,
    scan_price_template_catalog,
)


class V84PriceTemplateViewerTest(unittest.TestCase):
    def _write_price_file(
        self,
        price_dir: Path,
        name: str = "客户A-快递报价.xlsx",
        include_missing_column_sheet: bool = False,
    ) -> Path:
        workbook = openpyxl.Workbook()
        normal = workbook.active
        normal.title = "顺丰"
        normal.append(["省份参照列", "首重费用", "续重费用"])
        normal.append(["上海市", 6.5, 0.9])
        normal.append(["广东省", 6.9, 1.6])
        st = workbook.create_sheet("申通")
        st.append(["省份参照列", "首重费用", "续重费用"])
        st.append(["上海市", 3, 1])
        if include_missing_column_sheet:
            broken = workbook.create_sheet("缺列表")
            broken.append(["省份参照列", "首重费用"])
            broken.append(["上海市", 10])
        output_path = price_dir / name
        workbook.save(output_path)
        return output_path

    def test_scan_price_template_catalog_lists_customers_and_ignores_temp_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            price_dir = Path(temp_dir_text)
            first = self._write_price_file(price_dir, "客户B-快递报价.xlsx")
            second = self._write_price_file(price_dir, "客户A-快递报价.xlsx")
            self._write_price_file(price_dir, "~$客户C-快递报价.xlsx")

            catalog = scan_price_template_catalog(price_dir)

            self.assertEqual(catalog.customer_names, ["客户A", "客户B"])
            self.assertEqual(catalog.summaries[0].price_file, second)
            self.assertEqual(catalog.summaries[1].price_file, first)
            self.assertEqual(catalog.summaries[0].sheet_names, ["顺丰", "申通"])
            self.assertEqual(catalog.summaries[0].status, "正常")

    def test_scan_price_template_catalog_empty_directory_reports_clear_error(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            price_dir = Path(temp_dir_text)

            with self.assertRaisesRegex(FileNotFoundError, "目录中没有找到可用的 .xlsx 报价文件"):
                scan_price_template_catalog(price_dir)

    def test_load_price_template_workbook_keeps_sheet_order_and_rows(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            price_dir = Path(temp_dir_text)
            price_file = self._write_price_file(price_dir)

            workbook = load_price_template_workbook(price_file, customer="客户A")

            self.assertEqual(workbook.customer, "客户A")
            self.assertEqual(workbook.price_file, price_file)
            self.assertEqual([sheet.sheet_name for sheet in workbook.sheets], ["顺丰", "申通"])
            self.assertEqual(workbook.sheets[0].headers, ["省份参照列", "首重费用", "续重费用"])
            self.assertEqual(workbook.sheets[0].rows[0].province, "上海市")
            self.assertEqual(workbook.sheets[0].rows[0].first_price, 6.5)
            self.assertEqual(workbook.sheets[0].rows[0].extra_price, 0.9)
            self.assertEqual(workbook.sheets[0].rows[0].row_number, 2)

    def test_load_price_template_workbook_marks_sheet_missing_columns_without_failing(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            price_dir = Path(temp_dir_text)
            price_file = self._write_price_file(
                price_dir,
                include_missing_column_sheet=True,
            )

            workbook = load_price_template_workbook(price_file, customer="客户A")

            broken = workbook.sheets[-1]
            self.assertEqual(broken.sheet_name, "缺列表")
            self.assertEqual(broken.missing_columns, ["续重费用"])
            self.assertIn("缺少必要列：续重费用", broken.errors[0])
            self.assertEqual(broken.rows, [])

    def test_old_single_ticket_preview_api_is_removed(self) -> None:
        self.assertFalse(hasattr(calculator, "preview_express_fee"))
        self.assertFalse(hasattr(calculator, "PricePreviewResult"))


if __name__ == "__main__":
    unittest.main()
