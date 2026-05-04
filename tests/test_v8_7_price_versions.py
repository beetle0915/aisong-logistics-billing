from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from express_app.core.calculator import (  # noqa: E402
    CUSTOMER_HISTORY_DETAIL_SHEET,
    build_customer_history_summary,
    build_default_rule_config,
    customer_history_summary_file_name,
    run_express_fee_batch_job,
    scan_price_template_catalog,
    validate_express_fee_batch_job,
)
from express_app.core.models import ExpressFeeBatchJobConfig  # noqa: E402


class V87PriceVersionTest(unittest.TestCase):
    def _write_sales_file(self, path: Path, rows: list[list[object]]) -> None:
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.append(["出库单号", "出库日期", "业务员", "快递公司", "省", "重量"])
        for row in rows:
            ws.append(row)
        workbook.save(path)

    def _write_price_file(
        self,
        path: Path,
        sheets: dict[str, list[tuple[str, float, float]]] | None = None,
    ) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = openpyxl.Workbook()
        for index, (sheet_name, rows) in enumerate((sheets or {"顺丰": [("广东", 10, 2)]}).items()):
            ws = workbook.active if index == 0 else workbook.create_sheet()
            ws.title = sheet_name
            ws.append(["省份参照列", "首重费用", "续重费用"])
            for row in rows:
                ws.append(row)
        workbook.save(path)
        return path

    def _config(self, temp_dir: Path, sales_file: Path, price_dir: Path) -> ExpressFeeBatchJobConfig:
        return ExpressFeeBatchJobConfig(
            sales_files=[sales_file],
            price_dir=price_dir,
            output_dir=temp_dir / "output",
            split_dir=temp_dir / "split",
            split_customer_daily_files=True,
            generate_customer_history=True,
            rule_config=build_default_rule_config(),
        )

    def test_run_uses_latest_effective_salesman_price_version_by_shipping_date(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            price_dir = temp_dir / "prices"
            sales_file = temp_dir / "销售出库单.xlsx"
            self._write_sales_file(
                sales_file,
                [
                    ["CK001", date(2026, 4, 5), "客户A", "顺丰", "广东", 2],
                    ["CK002", date(2026, 4, 11), "客户A", "顺丰", "广东", 2],
                ],
            )
            self._write_price_file(
                price_dir / "客户A" / "20260401客户A-快递报价.xlsx",
                {"顺丰": [("广东", 10, 2)]},
            )
            self._write_price_file(
                price_dir / "客户A" / "20260410客户A-快递报价.xlsx",
                {"顺丰": [("广东", 20, 3)]},
            )

            result = run_express_fee_batch_job(self._config(temp_dir, sales_file, price_dir))

            self.assertTrue(result.ok)
            output_path = result.job_results[0].output_path
            workbook = openpyxl.load_workbook(output_path, data_only=True)
            ws = workbook.active
            headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
            header_map = {header: index + 1 for index, header in enumerate(headers)}
            self.assertEqual(ws.cell(row=2, column=header_map["快递费用"]).value, 12)
            self.assertEqual(ws.cell(row=2, column=header_map["报价版本"]).value, "20260401")
            self.assertEqual(ws.cell(row=2, column=header_map["报价生效日期"]).value, "2026-04-01")
            self.assertEqual(ws.cell(row=3, column=header_map["快递费用"]).value, 23)
            self.assertEqual(ws.cell(row=3, column=header_map["报价版本"]).value, "20260410")
            self.assertEqual(ws.cell(row=3, column=header_map["报价生效日期"]).value, "2026-04-10")

            detail_path = temp_dir / "split" / "客户A" / "2026-04-05_客户A_快递费明细.xlsx"
            detail_workbook = openpyxl.load_workbook(detail_path, data_only=True)
            detail_sheet = detail_workbook["快递明细"]
            detail_headers = [
                detail_sheet.cell(row=1, column=column).value
                for column in range(1, detail_sheet.max_column + 1)
            ]
            self.assertIn("报价版本", detail_headers)
            self.assertIn("报价生效日期", detail_headers)

    def test_legacy_top_level_price_file_still_works_with_blank_price_version_columns(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            sales_file = temp_dir / "销售出库单.xlsx"
            self._write_sales_file(
                sales_file,
                [["CK001", date(2026, 4, 5), "客户A", "顺丰", "广东", 2]],
            )
            self._write_price_file(price_dir / "客户A-快递报价.xlsx")

            result = run_express_fee_batch_job(self._config(temp_dir, sales_file, price_dir))

            self.assertTrue(result.ok)
            workbook = openpyxl.load_workbook(result.job_results[0].output_path, data_only=True)
            ws = workbook.active
            headers = [ws.cell(row=1, column=column).value for column in range(1, ws.max_column + 1)]
            header_map = {header: index + 1 for index, header in enumerate(headers)}
            self.assertEqual(ws.cell(row=2, column=header_map["快递费用"]).value, 12)
            self.assertIsNone(ws.cell(row=2, column=header_map["报价版本"]).value)
            self.assertIsNone(ws.cell(row=2, column=header_map["报价生效日期"]).value)

    def test_preflight_reports_missing_effective_price_version(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            price_dir = temp_dir / "prices"
            sales_file = temp_dir / "销售出库单.xlsx"
            self._write_sales_file(
                sales_file,
                [["CK001", date(2026, 4, 5), "客户A", "顺丰", "广东", 2]],
            )
            self._write_price_file(price_dir / "客户A" / "20260410客户A-快递报价.xlsx")

            result = validate_express_fee_batch_job(self._config(temp_dir, sales_file, price_dir))

            errors = "\n".join(result.errors)
            self.assertFalse(result.ok)
            self.assertIn("没有可用业务员报价版本", errors)
            self.assertIn("客户A", errors)
            self.assertIn("2026-04-05", errors)

    def test_preflight_reports_duplicate_salesman_price_versions(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            price_dir = temp_dir / "prices"
            sales_file = temp_dir / "销售出库单.xlsx"
            self._write_sales_file(
                sales_file,
                [["CK001", date(2026, 4, 5), "客户A", "顺丰", "广东", 2]],
            )
            self._write_price_file(price_dir / "客户A" / "20260401客户A-快递报价.xlsx")
            self._write_price_file(price_dir / "客户A" / "20260401客户A-快递报价备份.xlsx")

            result = validate_express_fee_batch_job(self._config(temp_dir, sales_file, price_dir))

            errors = "\n".join(result.errors)
            self.assertFalse(result.ok)
            self.assertIn("重复报价版本", errors)
            self.assertIn("客户A", errors)
            self.assertIn("2026-04-01", errors)

    def test_history_detail_merges_new_price_version_columns_and_keeps_old_rows_blank(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            customer_dir = Path(temp_dir_text) / "客户A"
            customer_dir.mkdir()
            self._write_daily_detail(
                customer_dir,
                ["出库单号", "出库日期", "业务员", "重量", "快递费用", "快递公司（标准版）"],
                ["CK001", date(2026, 4, 1), "客户A", 1, 10, "顺丰"],
            )
            build_customer_history_summary(customer_dir, build_default_rule_config())
            self._write_daily_detail(
                customer_dir,
                [
                    "出库单号",
                    "出库日期",
                    "业务员",
                    "重量",
                    "快递费用",
                    "快递公司（标准版）",
                    "报价版本",
                    "报价生效日期",
                ],
                ["CK002", date(2026, 4, 2), "客户A", 1, 12, "顺丰", "20260401", "2026-04-01"],
            )

            build_customer_history_summary(customer_dir, build_default_rule_config())

            workbook = openpyxl.load_workbook(
                customer_dir / customer_history_summary_file_name("客户A"),
                data_only=True,
            )
            detail_sheet = workbook[CUSTOMER_HISTORY_DETAIL_SHEET]
            headers = [
                detail_sheet.cell(row=1, column=column).value
                for column in range(1, detail_sheet.max_column + 1)
            ]
            version_col = headers.index("报价版本") + 1
            effective_col = headers.index("报价生效日期") + 1
            self.assertIsNone(detail_sheet.cell(row=2, column=version_col).value)
            self.assertIsNone(detail_sheet.cell(row=2, column=effective_col).value)
            self.assertEqual(detail_sheet.cell(row=3, column=version_col).value, "20260401")
            self.assertEqual(detail_sheet.cell(row=3, column=effective_col).value, "2026-04-01")

    def test_price_preview_catalog_lists_salesman_folder_versions(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            price_dir = Path(temp_dir_text)
            self._write_price_file(price_dir / "客户A" / "20260401客户A-快递报价.xlsx")
            self._write_price_file(price_dir / "客户A" / "20260410客户A-快递报价.xlsx")

            catalog = scan_price_template_catalog(price_dir)

            self.assertEqual(catalog.customer_names, ["客户A"])
            self.assertEqual(
                [summary.version for summary in catalog.summaries],
                ["20260401", "20260410"],
            )
            self.assertEqual(catalog.summaries[-1].price_file.name, "20260410客户A-快递报价.xlsx")

    def _write_daily_detail(self, customer_dir: Path, headers: list[str], row: list[object]) -> None:
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = "快递明细"
        ws.append(headers)
        ws.append(row)
        shipping_date = row[headers.index("出库日期")]
        workbook.save(customer_dir / f"{shipping_date:%Y-%m-%d}_客户A_快递费明细.xlsx")


if __name__ == "__main__":
    unittest.main()
