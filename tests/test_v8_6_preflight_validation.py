from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from express_app.core.calculator import (  # noqa: E402
    build_default_rule_config,
    run_express_fee_batch_job,
    validate_express_fee_batch_job,
)
from express_app.core.models import ExpressFeeBatchJobConfig  # noqa: E402


class V86PreflightValidationTest(unittest.TestCase):
    def _write_sales_file(self, path: Path, rows: list[list[object]]) -> None:
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.append(["出库单号", "出库日期", "业务员", "快递公司", "省", "重量"])
        for row in rows:
            ws.append(row)
        workbook.save(path)

    def _write_price_file(
        self,
        price_dir: Path,
        salesman: str = "客户A",
        sheets: dict[str, list[tuple[str, float, float]]] | None = None,
    ) -> Path:
        workbook = openpyxl.Workbook()
        default_sheets = {"顺丰": [("广东", 10, 2)]}
        for index, (sheet_name, rows) in enumerate((sheets or default_sheets).items()):
            ws = workbook.active if index == 0 else workbook.create_sheet()
            ws.title = sheet_name
            ws.append(["省份参照列", "首重费用", "续重费用"])
            for row in rows:
                ws.append(row)
        output_path = price_dir / f"{salesman}-快递报价.xlsx"
        workbook.save(output_path)
        return output_path

    def _config(self, temp_dir: Path, sales_file: Path, price_dir: Path) -> ExpressFeeBatchJobConfig:
        return ExpressFeeBatchJobConfig(
            sales_files=[sales_file],
            price_dir=price_dir,
            output_dir=temp_dir / "output",
            split_dir=temp_dir / "split",
            split_customer_daily_files=True,
            generate_customer_history=True,
            rule_config=build_default_rule_config(),
        )

    def test_preflight_validation_passes_without_writing_result_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            sales_file = temp_dir / "销售出库单.xlsx"
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            self._write_sales_file(
                sales_file,
                [["CK001", date(2026, 4, 7), "客户A", "顺丰", "广东", 2.5]],
            )
            self._write_price_file(price_dir)

            result = validate_express_fee_batch_job(self._config(temp_dir, sales_file, price_dir))

            self.assertTrue(result.ok)
            self.assertEqual(result.total_rows, 1)
            self.assertEqual(result.success_rows, 1)
            self.assertEqual(result.failed_rows, 0)
            self.assertEqual(result.errors, [])
            self.assertFalse((temp_dir / "output").exists())
            self.assertFalse((temp_dir / "split").exists())

    def test_preflight_validation_blocks_bad_input_without_polluting_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            sales_file = temp_dir / "销售出库单.xlsx"
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            self._write_sales_file(
                sales_file,
                [
                    ["CK001", date(2026, 4, 7), "客户A", "顺丰", "广东", "abc"],
                    ["CK002", date(2026, 4, 7), "客户B", "顺丰", "广东", 2],
                    ["CK003", date(2026, 4, 7), "客户A", "顺丰", "广东", 20],
                ],
            )
            self._write_price_file(price_dir)

            result = validate_express_fee_batch_job(self._config(temp_dir, sales_file, price_dir))

            errors = "\n".join(result.errors)
            self.assertFalse(result.ok)
            self.assertEqual(result.total_rows, 3)
            self.assertEqual(result.success_rows, 0)
            self.assertEqual(result.failed_rows, 3)
            self.assertIn("重量原值：abc", errors)
            self.assertIn("业务员没有价格表：客户B", errors)
            self.assertIn("缺少大件报价模板：顺丰_大件", errors)
            self.assertFalse((temp_dir / "output").exists())
            self.assertFalse((temp_dir / "split").exists())

    def test_preflight_validation_reports_non_positive_weight_before_price_lookup(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            sales_file = temp_dir / "销售出库单.xlsx"
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            self._write_sales_file(
                sales_file,
                [["CK001", date(2026, 4, 7), "客户A", "顺丰", "浙江", 0]],
            )
            self._write_price_file(price_dir)

            result = validate_express_fee_batch_job(self._config(temp_dir, sales_file, price_dir))

            errors = "\n".join(result.errors)
            self.assertFalse(result.ok)
            self.assertIn("重量必须大于0", errors)
            self.assertIn("请将重量改为大于 0 的数字", errors)
            self.assertNotIn("找不到对应报价", errors)

    def test_running_after_preflight_still_generates_result_files(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            sales_file = temp_dir / "销售出库单.xlsx"
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            self._write_sales_file(
                sales_file,
                [["CK001", date(2026, 4, 7), "客户A", "顺丰", "广东", 2.5]],
            )
            self._write_price_file(price_dir)
            config = self._config(temp_dir, sales_file, price_dir)

            validation = validate_express_fee_batch_job(config)
            run_result = run_express_fee_batch_job(config)

            self.assertTrue(validation.ok)
            self.assertTrue(run_result.ok)
            self.assertTrue((temp_dir / "output").exists())
            self.assertTrue((temp_dir / "split").exists())


if __name__ == "__main__":
    unittest.main()
