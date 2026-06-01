from __future__ import annotations

import json
import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from express_app.core.balance_upload import (  # noqa: E402
    build_balance_upload_payload,
    collect_balance_upload_preview,
    upload_balance_payload,
)
from express_app.core.calculator import (  # noqa: E402
    CUSTOMER_ABNORMAL_DEDUCTION_HEADERS,
    CUSTOMER_ABNORMAL_DEDUCTION_SHEET,
    CUSTOMER_HISTORY_SHEET,
    CUSTOMER_PAYMENT_HEADERS,
    CUSTOMER_PAYMENT_SHEET,
    customer_history_summary_file_name,
)


class _FakeResponse:
    status = 200

    def __enter__(self) -> "_FakeResponse":
        return self

    def __exit__(self, *_args: object) -> None:
        return None

    def read(self) -> bytes:
        return b'{"ok":true}'


class _FakeOpener:
    def __init__(self) -> None:
        self.request = None
        self.timeout = None

    def __call__(self, request, timeout: int):  # type: ignore[no-untyped-def]
        self.request = request
        self.timeout = timeout
        return _FakeResponse()


class BalanceUploadTest(unittest.TestCase):
    def _write_customer_history(
        self,
        split_dir: Path,
        customer: str,
        rows: list[tuple[date, float, float]],
        payments: list[tuple[date, float]] | None = None,
        abnormal_deductions: list[tuple[date, float]] | None = None,
        include_abnormal_sheet: bool = True,
    ) -> Path:
        customer_dir = split_dir / customer
        customer_dir.mkdir(parents=True)

        workbook = openpyxl.Workbook()
        history_sheet = workbook.active
        history_sheet.title = CUSTOMER_HISTORY_SHEET
        history_sheet.append(
            [
                "日期",
                "单数",
                "总重量",
                "今日快递总消费",
                "累计快递费用",
                "今日收款",
                "累计收款",
                "今日异常扣款",
                "累计异常扣款",
                "当前余额",
            ]
        )
        for row_index, (shipping_date, today_fee, cumulative_fee) in enumerate(rows, start=2):
            history_sheet.append(
                [
                    shipping_date,
                    1,
                    1.0,
                    today_fee,
                    cumulative_fee,
                    "",
                    "",
                    "",
                    "",
                    f"=G{row_index}-E{row_index}-I{row_index}",
                ]
            )

        payment_sheet = workbook.create_sheet(CUSTOMER_PAYMENT_SHEET)
        payment_sheet.append(CUSTOMER_PAYMENT_HEADERS)
        for paid_date, amount in payments or []:
            payment_sheet.append([paid_date, "微信", "财务", amount, paid_date])

        if include_abnormal_sheet:
            deduction_sheet = workbook.create_sheet(CUSTOMER_ABNORMAL_DEDUCTION_SHEET)
            deduction_sheet.append(CUSTOMER_ABNORMAL_DEDUCTION_HEADERS)
            for deduction_date, amount in abnormal_deductions or []:
                deduction_sheet.append([deduction_date, amount, "测试扣款"])

        output_path = customer_dir / customer_history_summary_file_name(customer)
        workbook.save(output_path)
        workbook.close()
        return output_path

    def test_collects_today_fee_and_balance_for_confirmed_upload_date(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            split_dir = Path(temp_dir_text)
            self._write_customer_history(
                split_dir,
                "客户A",
                [
                    (date(2026, 5, 8), 80.0, 80.0),
                    (date(2026, 5, 9), 120.0, 200.0),
                ],
                payments=[(date(2026, 5, 8), 500.0)],
                abnormal_deductions=[(date(2026, 5, 9), 30.0)],
            )
            self._write_customer_history(
                split_dir,
                "客户B",
                [(date(2026, 5, 9), 260.5, 260.5)],
                payments=[(date(2026, 5, 10), 999.0)],
                abnormal_deductions=[],
            )

            preview = collect_balance_upload_preview(split_dir, date(2026, 5, 9))

            self.assertEqual(preview.customer_count, 2)
            self.assertEqual(preview.total_today_fee, 380.5)
            self.assertEqual(preview.total_today_balance, 9.5)
            self.assertEqual(preview.debtor_count, 1)
            self.assertEqual([record.customer for record in preview.records], ["客户B", "客户A"])
            self.assertEqual(preview.records[0].today_fee, 260.5)
            self.assertEqual(preview.records[0].today_balance, -260.5)
            self.assertEqual(preview.records[0].balance_date, date(2026, 5, 9))
            self.assertEqual(preview.records[0].status, "欠款")
            self.assertEqual(preview.records[1].today_balance, 270.0)
            self.assertEqual(preview.records[1].balance_date, date(2026, 5, 9))

    def test_uses_latest_history_before_upload_date_when_customer_has_no_today_shipments(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            split_dir = Path(temp_dir_text)
            self._write_customer_history(
                split_dir,
                "客户A",
                [(date(2026, 5, 8), 80.0, 80.0)],
                payments=[(date(2026, 5, 8), 200.0), (date(2026, 5, 9), 50.0)],
                abnormal_deductions=[(date(2026, 5, 9), 20.0)],
            )

            preview = collect_balance_upload_preview(split_dir, date(2026, 5, 9))

            self.assertEqual(preview.errors, [])
            self.assertEqual(preview.customer_count, 1)
            self.assertEqual(preview.carried_forward_count, 1)
            self.assertEqual(preview.records[0].customer, "客户A")
            self.assertEqual(preview.records[0].today_fee, 0.0)
            self.assertEqual(preview.records[0].today_balance, 150.0)
            self.assertEqual(preview.records[0].balance_date, date(2026, 5, 8))

    def test_builds_minimal_payload_for_remote_upload(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            split_dir = Path(temp_dir_text)
            self._write_customer_history(
                split_dir,
                "客户A",
                [(date(2026, 5, 9), 120.0, 200.0)],
                payments=[(date(2026, 5, 9), 500.0)],
            )
            preview = collect_balance_upload_preview(split_dir, date(2026, 5, 9))

            payload = build_balance_upload_payload(
                preview,
                app_version="8.10.0",
                uploaded_at=datetime(2026, 5, 9, 12, 30, 0),
            )

            self.assertEqual(payload["schemaVersion"], "1.0")
            self.assertEqual(payload["appVersion"], "8.10.0")
            self.assertEqual(payload["uploadDate"], "2026-05-09")
            self.assertEqual(payload["uploadedAt"], "2026-05-09T12:30:00")
            self.assertEqual(
                payload["records"],
                [
                    {
                        "customer": "客户A",
                        "todayFee": 120.0,
                        "todayBalance": 300.0,
                        "balanceDate": "2026-05-09",
                        "status": "充足",
                    }
                ],
            )
            self.assertEqual(payload["summary"]["customerCount"], 1)

    def test_upload_posts_json_with_bearer_token(self) -> None:
        opener = _FakeOpener()

        result = upload_balance_payload(
            "https://example.test/upload",
            "secret-token",
            {"records": []},
            timeout=7,
            opener=opener,
        )

        self.assertTrue(result.ok)
        self.assertEqual(result.status_code, 200)
        self.assertEqual(opener.timeout, 7)
        self.assertIsNotNone(opener.request)
        self.assertEqual(opener.request.full_url, "https://example.test/upload")
        self.assertEqual(opener.request.get_method(), "POST")
        self.assertEqual(opener.request.headers["Authorization"], "Bearer secret-token")
        self.assertEqual(opener.request.headers["Content-type"], "application/json")
        self.assertEqual(json.loads(opener.request.data.decode("utf-8")), {"records": []})


if __name__ == "__main__":
    unittest.main()
