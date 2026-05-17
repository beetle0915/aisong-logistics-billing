from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from express_app.core.calculator import build_default_rule_config, run_express_fee_batch_job  # noqa: E402
from express_app.core.calculator import build_versioned_price_catalog  # noqa: E402
from express_app.core.models import ExpressFeeBatchJobConfig  # noqa: E402


class V891WeightTierTest(unittest.TestCase):
    def _write_sales_file(self, path: Path, weight: float) -> None:
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.append(["出库日期", "业务员", "快递公司", "省", "重量"])
        ws.append([date(2026, 5, 17), "客户A", "顺丰", "广东", weight])
        workbook.save(path)

    def _write_price_file(self, price_dir: Path) -> None:
        workbook = openpyxl.Workbook()
        sheets = {
            "顺丰": [("广东", 10, 2)],
            "顺丰_大件": [("广东", 30, 5)],
            "顺丰_超大件": [("广东", 80, 8)],
        }
        for index, (sheet_name, rows) in enumerate(sheets.items()):
            ws = workbook.active if index == 0 else workbook.create_sheet()
            ws.title = sheet_name
            ws.append(["省份参照列", "首重费用", "续重费用"])
            for row in rows:
                ws.append(row)
        workbook.save(price_dir / "客户A-快递报价.xlsx")

    def _run_job(self, temp_dir: Path, sales_file: Path, price_dir: Path):
        return run_express_fee_batch_job(
            ExpressFeeBatchJobConfig(
                sales_files=[sales_file],
                price_dir=price_dir,
                output_dir=temp_dir / "output",
                split_dir=temp_dir / "split",
                split_customer_daily_files=False,
                generate_customer_history=False,
                rule_config=build_default_rule_config(),
            )
        )

    def _read_result_row(self, output_path: Path) -> dict[str, object]:
        workbook = openpyxl.load_workbook(output_path, data_only=False)
        ws = workbook.active
        headers = [cell.value for cell in ws[1]]
        return {header: ws.cell(row=2, column=index + 1).value for index, header in enumerate(headers)}

    def test_large_piece_uses_large_threshold_as_first_weight(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            sales_file = temp_dir / "销售出库单.xlsx"
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            self._write_sales_file(sales_file, weight=22)
            self._write_price_file(price_dir)

            result = self._run_job(temp_dir, sales_file, price_dir)
            row = self._read_result_row(result.job_results[0].output_path)

            self.assertTrue(result.ok, "\n".join(result.logs))
            self.assertEqual(row["计费档位"], "大件")
            self.assertEqual(row["计费模板"], "顺丰_大件")
            self.assertEqual(row["首重重量"], 20)
            self.assertEqual(row["续重重量"], 2)
            self.assertEqual(row["快递费用"], 40)

    def test_super_large_piece_uses_60kg_threshold_as_first_weight(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            sales_file = temp_dir / "销售出库单.xlsx"
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            self._write_sales_file(sales_file, weight=65)
            self._write_price_file(price_dir)

            result = self._run_job(temp_dir, sales_file, price_dir)
            row = self._read_result_row(result.job_results[0].output_path)

            self.assertTrue(result.ok, "\n".join(result.logs))
            self.assertEqual(row["计费档位"], "超大件")
            self.assertEqual(row["计费模板"], "顺丰_超大件")
            self.assertEqual(row["首重重量"], 60)
            self.assertEqual(row["续重重量"], 5)
            self.assertEqual(row["快递费用"], 120)

    def test_missing_super_large_template_reports_required_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            sales_file = temp_dir / "销售出库单.xlsx"
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            self._write_sales_file(sales_file, weight=65)

            workbook = openpyxl.Workbook()
            ws = workbook.active
            ws.title = "顺丰"
            ws.append(["省份参照列", "首重费用", "续重费用"])
            ws.append(["广东", 10, 2])
            large_ws = workbook.create_sheet("顺丰_大件")
            large_ws.append(["省份参照列", "首重费用", "续重费用"])
            large_ws.append(["广东", 30, 5])
            workbook.save(price_dir / "客户A-快递报价.xlsx")

            result = self._run_job(temp_dir, sales_file, price_dir)

            logs = "\n".join(result.logs)
            self.assertFalse(result.ok)
            self.assertIn("缺少超大件报价模板：顺丰_超大件", logs)
            self.assertIn("请在 客户A 的报价表中新增 sheet「顺丰_超大件」", logs)

    def test_super_large_template_is_not_treated_as_standard_company(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir_text:
            temp_dir = Path(temp_dir_text)
            price_dir = temp_dir / "prices"
            price_dir.mkdir()
            self._write_price_file(price_dir)

            catalog = build_versioned_price_catalog(price_dir, rule_config=build_default_rule_config())

            self.assertIn("顺丰", catalog.standard_companies)
            self.assertNotIn("顺丰_大件", catalog.standard_companies)
            self.assertNotIn("顺丰_超大件", catalog.standard_companies)


if __name__ == "__main__":
    unittest.main()
