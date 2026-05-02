from __future__ import annotations

import sys
import tempfile
import unittest
from datetime import date, datetime
from pathlib import Path

import openpyxl

APP_DIR = Path(__file__).resolve().parents[1] / "app"
sys.path.insert(0, str(APP_DIR))

from express_app.core.calculator import (  # noqa: E402
    CUSTOMER_HISTORY_SHEET,
    CUSTOMER_HISTORY_SUMMARY_FILE,
    customer_history_summary_file_name,
    build_customer_history_summary,
    build_default_rule_config,
)


HISTORY_DETAIL_SHEET = "快递明细"
PAYMENT_SHEET = "收款记录"
LEGACY_HISTORY_FILE = "客户快递费历史汇总.xlsx"


class CustomerHistoryWorkbookTest(unittest.TestCase):
    def assert_cell_date(self, value: object, expected: date) -> None:
        if isinstance(value, datetime):
            value = value.date()
        self.assertEqual(value, expected)

    def assert_cell_has_thin_border(self, cell) -> None:
        self.assertEqual(cell.border.left.style, "thin")
        self.assertEqual(cell.border.right.style, "thin")
        self.assertEqual(cell.border.top.style, "thin")
        self.assertEqual(cell.border.bottom.style, "thin")

    def _write_daily_detail(
        self,
        customer_dir: Path,
        customer: str,
        shipping_date: date,
        fee: float,
        outbound_no: str = "CK001",
        weight: float = 2.5,
    ) -> None:
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = "快递明细"
        ws.append(["出库单号", "出库日期", "业务员", "重量", "快递费用", "快递公司（标准版）"])
        ws.append([outbound_no, shipping_date, customer, weight, fee, "顺丰"])
        output_path = customer_dir / f"{shipping_date:%Y-%m-%d}_{customer}_快递费明细.xlsx"
        workbook.save(output_path)

    def test_history_workbook_file_name_includes_customer_name(self) -> None:
        self.assertEqual(customer_history_summary_file_name("客户A"), "客户A_客户快递费历史汇总.xlsx")

        with tempfile.TemporaryDirectory() as temp_dir:
            customer_dir = Path(temp_dir) / "客户A"
            customer_dir.mkdir()
            self._write_daily_detail(customer_dir, "客户A", date(2026, 4, 2), 120.5)

            summary = build_customer_history_summary(customer_dir, build_default_rule_config())

            self.assertEqual(summary.output_path, customer_dir / "客户A_客户快递费历史汇总.xlsx")
            self.assertTrue(summary.output_path.exists())

    def test_history_workbook_adds_payment_sheet_and_balance_formulas(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            customer_dir = Path(temp_dir) / "客户A"
            customer_dir.mkdir()
            self._write_daily_detail(customer_dir, "客户A", date(2026, 4, 2), 120.5)

            build_customer_history_summary(customer_dir, build_default_rule_config())

            workbook = openpyxl.load_workbook(
                customer_dir / customer_history_summary_file_name("客户A"),
                data_only=False,
            )
            self.assertEqual(workbook.sheetnames, [CUSTOMER_HISTORY_SHEET, HISTORY_DETAIL_SHEET, PAYMENT_SHEET])

            history_sheet = workbook[CUSTOMER_HISTORY_SHEET]
            self.assertEqual(
                [history_sheet.cell(row=1, column=column).value for column in range(1, 13)],
                [
                    "日期",
                    "单数",
                    "总重量",
                    "今日快递总消费",
                    "累计快递费用",
                    "今日收款",
                    "累计收款",
                    "当前余额",
                    "顺丰单数",
                    "申通单数",
                    "德邦单数",
                    "大件单数",
                ],
            )
            self.assertIsNone(history_sheet.cell(row=1, column=13).value)
            self.assertGreaterEqual(history_sheet.row_dimensions[1].height, 28)
            self.assertGreaterEqual(history_sheet.row_dimensions[2].height, 26)
            self.assertGreaterEqual(history_sheet.column_dimensions["D"].width, 18)
            self.assertGreaterEqual(history_sheet.column_dimensions["H"].width, 16)
            self.assert_cell_has_thin_border(history_sheet["A1"])
            self.assert_cell_has_thin_border(history_sheet["H2"])
            self.assertEqual(
                history_sheet["F2"].value,
                '=SUMIFS(\'收款记录\'!$D:$D,\'收款记录\'!$A:$A,">="&A2,'
                '\'收款记录\'!$A:$A,"<"&A2+1)',
            )
            self.assertEqual(
                history_sheet["G2"].value,
                '=SUMIFS(\'收款记录\'!$D:$D,\'收款记录\'!$A:$A,"<"&A2+1,'
                '\'收款记录\'!$A:$A,"<>")',
            )
            self.assertEqual(history_sheet["H2"].value, "=G2-E2")

            detail_sheet = workbook[HISTORY_DETAIL_SHEET]
            self.assertEqual(
                [detail_sheet.cell(row=1, column=column).value for column in range(1, 7)],
                ["出库单号", "出库日期", "业务员", "重量", "快递费用", "快递公司（标准版）"],
            )
            self.assertEqual(detail_sheet["A2"].value, "CK001")
            self.assertEqual(detail_sheet["E2"].value, 120.5)
            self.assertEqual(detail_sheet["G1"].value, "历史记录Key")
            self.assertTrue(detail_sheet.column_dimensions["G"].hidden)
            self.assertEqual(detail_sheet["G2"].value, "CK001|2026-04-02 00:00:00")
            self.assertEqual(detail_sheet["I1"].value, "最后更新时间")
            self.assert_cell_has_thin_border(detail_sheet["A1"])
            self.assert_cell_has_thin_border(detail_sheet["I2"])

            payment_sheet = workbook[PAYMENT_SHEET]
            self.assertEqual(
                [payment_sheet.cell(row=1, column=column).value for column in range(1, 6)],
                ["支付时间", "支付方式", "收款人", "已付金额", "收款时间"],
            )
            self.assert_cell_has_thin_border(payment_sheet["A1"])
            self.assert_cell_has_thin_border(payment_sheet["E20"])

    def test_payment_records_are_preserved_when_history_is_refreshed(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            customer_dir = Path(temp_dir) / "客户A"
            customer_dir.mkdir()
            self._write_daily_detail(customer_dir, "客户A", date(2026, 4, 2), 120.5)

            workbook = openpyxl.Workbook()
            ws = workbook.active
            ws.title = PAYMENT_SHEET
            ws.append(["支付时间", "支付方式", "收款人", "已付金额", "收款时间"])
            ws.append([date(2026, 4, 2), "微信", "小李", 5000, date(2026, 4, 2)])
            workbook.save(customer_dir / CUSTOMER_HISTORY_SUMMARY_FILE)

            build_customer_history_summary(customer_dir, build_default_rule_config())

            refreshed = openpyxl.load_workbook(
                customer_dir / customer_history_summary_file_name("客户A"),
                data_only=False,
            )
            payment_sheet = refreshed[PAYMENT_SHEET]
            self.assert_cell_date(payment_sheet["A2"].value, date(2026, 4, 2))
            self.assertEqual(payment_sheet["B2"].value, "微信")
            self.assertEqual(payment_sheet["C2"].value, "小李")
            self.assertEqual(payment_sheet["D2"].value, 5000)
            self.assert_cell_date(payment_sheet["E2"].value, date(2026, 4, 2))

    def test_history_detail_upserts_by_outbound_number_and_shipping_time(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            customer_dir = Path(temp_dir) / "客户A"
            customer_dir.mkdir()
            self._write_daily_detail(
                customer_dir,
                "客户A",
                date(2026, 4, 2),
                120.5,
                outbound_no="CK001",
                weight=2.5,
            )
            build_customer_history_summary(customer_dir, build_default_rule_config())

            self._write_daily_detail(
                customer_dir,
                "客户A",
                date(2026, 4, 2),
                130.0,
                outbound_no="CK001",
                weight=3.0,
            )
            self._write_daily_detail(
                customer_dir,
                "客户A",
                date(2026, 4, 3),
                88.0,
                outbound_no="CK002",
                weight=1.5,
            )
            build_customer_history_summary(customer_dir, build_default_rule_config())

            refreshed = openpyxl.load_workbook(
                customer_dir / customer_history_summary_file_name("客户A"),
                data_only=False,
            )
            detail_sheet = refreshed[HISTORY_DETAIL_SHEET]
            self.assertEqual(detail_sheet.max_row, 3)
            self.assertEqual(detail_sheet["A2"].value, "CK001")
            self.assertEqual(detail_sheet["D2"].value, 3.0)
            self.assertEqual(detail_sheet["E2"].value, 130.0)
            self.assertEqual(detail_sheet["A3"].value, "CK002")
            self.assertEqual(detail_sheet["E3"].value, 88.0)
            self.assertNotEqual(detail_sheet["H2"].value, None)
            self.assertNotEqual(detail_sheet["I2"].value, None)


if __name__ == "__main__":
    unittest.main()
