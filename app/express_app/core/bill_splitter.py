from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


BILL_DETAIL_SHEET_NAME = "账单明细"
DEFAULT_SPLIT_FIELD = "经手人"
TRACKING_NUMBER_FIELD = "运单号码"
PAYABLE_AMOUNT_FIELD = "应付金额"
SOURCE_FILE_FIELD = "来源文件"
SOURCE_ROW_FIELD = "来源行号"

INVALID_FILE_CHARS_RE = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')
INVALID_SHEET_CHARS_RE = re.compile(r"[\[\]\\/:*?]+")
SUMMARY_LABEL_PREFIXES = ("合计", "小计", "总计", "总额", "汇总")
INVALID_SPLIT_VALUES = {"#N/A", "N/A", "无"}


@dataclass(frozen=True)
class BillSplitOutput:
    split_value: str
    output_path: Path
    row_count: int


@dataclass(frozen=True)
class BillSplitResult:
    output_dir: Path
    common_headers: list[str]
    output_paths: list[Path]
    outputs: list[BillSplitOutput]


@dataclass(frozen=True)
class BillSplitFileScan:
    path: Path
    sheet_name: str
    headers: list[str]
    total_rows: int
    error: str = ""


@dataclass(frozen=True)
class BillSplitScanResult:
    output_dir: Path
    common_headers: list[str]
    files: list[BillSplitFileScan]
    errors: list[str]
    logs: list[str]


@dataclass(frozen=True)
class BillSource:
    path: Path
    headers: list[str]
    header_indexes: dict[str, int]
    rows: list[tuple[int, tuple[Any, ...]]]


@dataclass(frozen=True)
class OutputRow:
    source_path: Path
    source_row_number: int
    values: list[Any]


@dataclass(frozen=True)
class BillScanContext:
    result: BillSplitScanResult
    sources: list[BillSource]


def scan_bill_split_directory(input_dir: Path | str, split_field: str = DEFAULT_SPLIT_FIELD) -> BillSplitScanResult:
    return _scan_bill_split_directory(Path(input_dir), split_field).result


def split_bills_by_field(input_dir: Path | str, split_field: str = DEFAULT_SPLIT_FIELD) -> BillSplitResult:
    input_path = Path(input_dir)
    scan_context = _scan_bill_split_directory(input_path, split_field)
    scan_result = scan_context.result

    sources = scan_context.sources
    common_headers = scan_result.common_headers
    if not sources:
        scan_result.output_dir.mkdir(parents=True, exist_ok=True)
        return BillSplitResult(output_dir=scan_result.output_dir, common_headers=[], output_paths=[], outputs=[])
    if split_field not in common_headers:
        raise ValueError(f"共同表头缺少拆分字段：{split_field}")
    if TRACKING_NUMBER_FIELD not in common_headers:
        raise ValueError(f"共同表头缺少必需字段：{TRACKING_NUMBER_FIELD}")

    rows_by_split_value: dict[str, list[OutputRow]] = {}
    split_file_value_by_key: dict[str, Any] = {}
    for source in sources:
        for source_row_number, row in source.rows:
            if not _is_valid_data_row(row, source.header_indexes, split_field):
                continue
            split_value = row[source.header_indexes[split_field]]
            split_key = _display_value(split_value)
            values = [_row_value(row, source.header_indexes[header]) for header in common_headers]
            rows_by_split_value.setdefault(split_key, []).append(
                OutputRow(source_path=source.path, source_row_number=source_row_number, values=values)
            )
            split_file_value_by_key.setdefault(split_key, split_value)

    scan_result.output_dir.mkdir(parents=True, exist_ok=True)
    outputs: list[BillSplitOutput] = []
    used_output_names: set[str] = set()
    for split_key in sorted(rows_by_split_value):
        output_rows = rows_by_split_value[split_key]
        output_file_name = _unique_output_file_name(
            _build_output_file_name(split_field, split_file_value_by_key[split_key]),
            used_output_names,
        )
        output_path = scan_result.output_dir / output_file_name
        _write_split_workbook(output_path, common_headers, output_rows)
        outputs.append(
            BillSplitOutput(
                split_value=split_key,
                output_path=output_path,
                row_count=len(output_rows),
            )
        )

    return BillSplitResult(
        output_dir=scan_result.output_dir,
        common_headers=common_headers,
        output_paths=[output.output_path for output in outputs],
        outputs=outputs,
    )


def _scan_bill_split_directory(input_dir: Path, split_field: str) -> BillScanContext:
    output_dir = input_dir.parent / _build_output_dir_name(input_dir.name, split_field)
    sources, file_scans, logs = _load_bill_sources(input_dir)
    common_headers = _common_headers(sources)
    errors: list[str] = []
    if sources and split_field not in common_headers:
        errors.append(f"共同表头缺少拆分字段：{split_field}")
    if sources and TRACKING_NUMBER_FIELD not in common_headers:
        errors.append(f"共同表头缺少必需字段：{TRACKING_NUMBER_FIELD}")

    result = BillSplitScanResult(
        output_dir=output_dir,
        common_headers=common_headers,
        files=file_scans,
        errors=errors,
        logs=logs,
    )
    return BillScanContext(result=result, sources=sources)


def build_bill_split_output_dir(input_dir: Path | str, split_field: str = DEFAULT_SPLIT_FIELD) -> Path:
    input_path = Path(input_dir)
    return input_path.parent / _build_output_dir_name(input_path.name, split_field)


def _build_output_dir_name(input_name: str, split_field: str) -> str:
    field_name = _sanitize_file_part(split_field or DEFAULT_SPLIT_FIELD, max_length=80)
    return f"{input_name}拆分结果【拆分字段：{field_name}】"


def _load_bill_sources(input_dir: Path) -> tuple[list[BillSource], list[BillSplitFileScan], list[str]]:
    sources: list[BillSource] = []
    file_scans: list[BillSplitFileScan] = []
    logs: list[str] = []
    for path in sorted(input_dir.iterdir(), key=lambda item: item.name):
        if not _is_candidate_workbook(path):
            if path.is_file():
                logs.append(f"跳过非目标文件：{path.name}")
            continue

        try:
            workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
        except Exception as exc:
            error = f"读取失败：{exc}"
            file_scans.append(
                BillSplitFileScan(
                    path=path,
                    sheet_name=BILL_DETAIL_SHEET_NAME,
                    headers=[],
                    total_rows=0,
                    error=error,
                )
            )
            logs.append(f"跳过读取失败的文件：{path.name}，{exc}")
            continue
        try:
            if BILL_DETAIL_SHEET_NAME not in workbook.sheetnames:
                file_scans.append(
                    BillSplitFileScan(
                        path=path,
                        sheet_name=BILL_DETAIL_SHEET_NAME,
                        headers=[],
                        total_rows=0,
                        error=f"缺少{BILL_DETAIL_SHEET_NAME}",
                    )
                )
                logs.append(f"跳过缺少账单明细的文件：{path.name}")
                continue
            ws = workbook[BILL_DETAIL_SHEET_NAME]
            raw_headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
            headers, header_indexes = _build_header_indexes(raw_headers)
            if not headers:
                file_scans.append(
                    BillSplitFileScan(
                        path=path,
                        sheet_name=BILL_DETAIL_SHEET_NAME,
                        headers=[],
                        total_rows=0,
                        error="表头为空",
                    )
                )
                logs.append(f"跳过表头为空的文件：{path.name}")
                continue
            rows = [
                (row_number, tuple(row))
                for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2)
                if not _is_empty_row(tuple(row))
            ]
            sources.append(BillSource(path=path, headers=headers, header_indexes=header_indexes, rows=rows))
            file_scans.append(
                BillSplitFileScan(
                    path=path,
                    sheet_name=BILL_DETAIL_SHEET_NAME,
                    headers=headers,
                    total_rows=len(rows),
                )
            )
            logs.append(f"识别账单明细文件：{path.name}")
        finally:
            workbook.close()
    return sources, file_scans, logs


def _is_candidate_workbook(path: Path) -> bool:
    return (
        path.is_file()
        and path.suffix.lower() == ".xlsx"
        and not path.name.startswith(".")
        and not path.name.startswith("~$")
    )


def _build_header_indexes(raw_headers: tuple[Any, ...]) -> tuple[list[str], dict[str, int]]:
    headers: list[str] = []
    header_indexes: dict[str, int] = {}
    for index, value in enumerate(raw_headers):
        header = _display_value(value)
        if not header or header in header_indexes:
            continue
        headers.append(header)
        header_indexes[header] = index
    return headers, header_indexes


def _common_headers(sources: list[BillSource]) -> list[str]:
    if not sources:
        return []
    common = set(sources[0].headers)
    for source in sources[1:]:
        common &= set(source.headers)
    return [header for header in sources[0].headers if header in common]


def _is_valid_data_row(row: tuple[Any, ...], header_indexes: dict[str, int], split_field: str) -> bool:
    if _is_empty_row(row):
        return False
    first_value = _row_value(row, 0)
    if _is_summary_label(first_value):
        return False

    tracking_number = _row_value(row, header_indexes[TRACKING_NUMBER_FIELD])
    if _is_blank(tracking_number):
        return False

    split_value = _row_value(row, header_indexes[split_field])
    if _is_invalid_split_value(split_value):
        return False
    return True


def _row_value(row: tuple[Any, ...], index: int) -> Any:
    if index >= len(row):
        return None
    return row[index]


def _is_empty_row(row: tuple[Any, ...]) -> bool:
    return all(_is_blank(value) for value in row)


def _is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def _is_summary_label(value: Any) -> bool:
    if value is None:
        return False
    label = _display_value(value)
    normalized = re.sub(r"\s+", "", label).rstrip(":：")
    return normalized.startswith(SUMMARY_LABEL_PREFIXES)


def _is_invalid_split_value(value: Any) -> bool:
    if _is_blank(value):
        return True
    normalized = _display_value(value).upper()
    return normalized in INVALID_SPLIT_VALUES


def _build_output_file_name(split_field: str, split_value: Any) -> str:
    field_name = _sanitize_file_part(split_field, max_length=40)
    value_name = _sanitize_file_part(_display_value(split_value), max_length=120)
    return f"{field_name}_{value_name}.xlsx"


def _unique_output_file_name(file_name: str, used_names: set[str]) -> str:
    if file_name not in used_names:
        used_names.add(file_name)
        return file_name

    stem = file_name[:-5] if file_name.lower().endswith(".xlsx") else file_name
    suffix_number = 2
    while True:
        suffix = f"_{suffix_number}"
        candidate = f"{stem[: 180 - len(suffix)].rstrip(' ._')}{suffix}.xlsx"
        if candidate not in used_names:
            used_names.add(candidate)
            return candidate
        suffix_number += 1


def _sanitize_file_part(value: Any, max_length: int) -> str:
    sanitized = INVALID_FILE_CHARS_RE.sub("_", _display_value(value))
    sanitized = re.sub(r"\s+", " ", sanitized).strip(" .")
    sanitized = re.sub(r"_+", "_", sanitized).strip("_")
    if not sanitized:
        sanitized = "未命名"
    return sanitized[:max_length].rstrip(" ._") or "未命名"


def _sanitize_sheet_name(stem: str, used_names: set[str]) -> str:
    base = INVALID_SHEET_CHARS_RE.sub("_", stem).strip()
    base = re.sub(r"_+", "_", base).strip("'") or "Sheet"
    base = base[:31].rstrip() or "Sheet"
    candidate = base
    suffix_number = 2
    while candidate in used_names:
        suffix = f"_{suffix_number}"
        candidate = f"{base[: 31 - len(suffix)].rstrip()}{suffix}"
        suffix_number += 1
    used_names.add(candidate)
    return candidate


def _write_split_workbook(output_path: Path, common_headers: list[str], output_rows: list[OutputRow]) -> None:
    workbook = openpyxl.Workbook()
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_sheet_names: set[str] = set()
    rows_by_source: dict[Path, list[OutputRow]] = {}
    for output_row in output_rows:
        rows_by_source.setdefault(output_row.source_path, []).append(output_row)

    headers = common_headers + [SOURCE_FILE_FIELD, SOURCE_ROW_FIELD]
    for source_path in sorted(rows_by_source, key=lambda path: path.name):
        sheet_name = _sanitize_sheet_name(source_path.stem, used_sheet_names)
        ws = workbook.create_sheet(sheet_name)
        source_rows = rows_by_source[source_path]
        ws.append(_build_payable_summary_row(headers, source_rows))
        ws.append(headers)
        for output_row in source_rows:
            ws.append(output_row.values + [source_path.name, output_row.source_row_number])
        _style_output_sheet(ws)

    workbook.save(output_path)


def _build_payable_summary_row(headers: list[str], output_rows: list[OutputRow]) -> list[str]:
    summary_row = [""] * len(headers)
    summary_row[0] = "账款金额"
    if PAYABLE_AMOUNT_FIELD in headers:
        amount_column = get_column_letter(headers.index(PAYABLE_AMOUNT_FIELD) + 1)
        first_data_row = 3
        last_data_row = first_data_row + len(output_rows) - 1
        summary_row[1] = f"=SUM({amount_column}{first_data_row}:{amount_column}{last_data_row})"
    else:
        summary_row[1] = f"缺少{PAYABLE_AMOUNT_FIELD}字段"
    return summary_row


def _style_output_sheet(ws: openpyxl.worksheet.worksheet.Worksheet) -> None:
    summary_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True, size=16)
        cell.fill = summary_fill
    ws.row_dimensions[1].height = 30
    ws["B1"].number_format = '#,##0.00'

    for cell in ws[2]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A3"
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A2:{get_column_letter(ws.max_column)}{ws.max_row}"
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_length = max(len(_display_value(cell.value)) for cell in column_cells)
        ws.column_dimensions[letter].width = min(max(max_length + 2, 10), 32)


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()
