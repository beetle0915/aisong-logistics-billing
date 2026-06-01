"""Prepare and upload daily customer balance snapshots."""

from __future__ import annotations

import json
import urllib.error
import urllib.request
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

import openpyxl

from express_app.core.account_balance import _parse_date, _parse_number, _read_headers
from express_app.core.calculator import (
    CUSTOMER_ABNORMAL_DEDUCTION_SHEET,
    CUSTOMER_HISTORY_SHEET,
    CUSTOMER_PAYMENT_SHEET,
    customer_history_summary_file_name,
)


@dataclass(frozen=True)
class BalanceUploadRecord:
    customer: str
    today_fee: float
    today_balance: float
    balance_date: date
    status: str
    history_file: Path


@dataclass(frozen=True)
class BalanceUploadPreview:
    upload_date: date
    source_dir: Path
    records: list[BalanceUploadRecord] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def customer_count(self) -> int:
        return len(self.records)

    @property
    def total_today_fee(self) -> float:
        return round(sum(record.today_fee for record in self.records), 2)

    @property
    def total_today_balance(self) -> float:
        return round(sum(record.today_balance for record in self.records), 2)

    @property
    def debtor_count(self) -> int:
        return sum(1 for record in self.records if record.today_balance < 0)

    @property
    def carried_forward_count(self) -> int:
        return sum(1 for record in self.records if record.balance_date != self.upload_date)


@dataclass(frozen=True)
class BalanceUploadResult:
    ok: bool
    status_code: int | None
    message: str


def collect_balance_upload_preview(source_dir: Path, upload_date: date) -> BalanceUploadPreview:
    if not source_dir.exists():
        return BalanceUploadPreview(
            upload_date=upload_date,
            source_dir=source_dir,
            errors=[f"客户明细目录不存在：{source_dir}"],
        )

    records: list[BalanceUploadRecord] = []
    errors: list[str] = []
    for customer_dir in sorted(path for path in source_dir.iterdir() if path.is_dir()):
        history_file = customer_dir / customer_history_summary_file_name(customer_dir.name)
        if not history_file.exists():
            continue
        try:
            records.append(read_balance_upload_record(customer_dir.name, history_file, upload_date))
        except ValueError as exc:
            errors.append(f"{customer_dir.name}：{exc}")

    records.sort(key=lambda item: (item.today_balance >= 0, item.today_balance, item.customer))
    return BalanceUploadPreview(upload_date=upload_date, source_dir=source_dir, records=records, errors=errors)


def read_balance_upload_record(customer: str, history_file: Path, upload_date: date) -> BalanceUploadRecord:
    workbook = openpyxl.load_workbook(history_file, data_only=False, read_only=True)
    try:
        if CUSTOMER_HISTORY_SHEET not in workbook.sheetnames:
            raise ValueError(f"缺少 sheet：{CUSTOMER_HISTORY_SHEET}")

        history_sheet = workbook[CUSTOMER_HISTORY_SHEET]
        headers = _read_headers(history_sheet)
        required_headers = ["日期", "今日快递总消费", "累计快递费用"]
        missing_headers = [header for header in required_headers if header not in headers]
        if missing_headers:
            raise ValueError(f"历史汇总缺少列：{'、'.join(missing_headers)}")

        target_row: tuple[Any, ...] | None = None
        latest_row: tuple[Any, ...] | None = None
        latest_date: date | None = None
        date_index = headers["日期"]
        for row in history_sheet.iter_rows(min_row=2, values_only=True):
            row_date = _parse_date(row[date_index])
            if row_date is None:
                continue
            if row_date == upload_date:
                target_row = row
                latest_row = row
                latest_date = row_date
                continue
            if row_date < upload_date and (latest_date is None or row_date > latest_date):
                latest_row = row
                latest_date = row_date

        if latest_row is None or latest_date is None:
            raise ValueError(f"历史汇总没有 {upload_date.isoformat()} 之前的记录")

        today_fee = _parse_number(target_row[headers["今日快递总消费"]]) if target_row else 0.0
        cumulative_fee = _parse_number(latest_row[headers["累计快递费用"]])
        if today_fee is None:
            raise ValueError(f"{upload_date.isoformat()} 今日快递总消费为空")
        if cumulative_fee is None:
            raise ValueError(f"{latest_date.isoformat()} 累计快递费用为空")

        current_balance = _parse_number(
            latest_row[headers["当前余额"]] if "当前余额" in headers else None
        )
        if current_balance is None:
            total_paid = _sum_records_until(workbook, CUSTOMER_PAYMENT_SHEET, date_column=0, amount_column=3, end_date=upload_date)
            total_abnormal_deducted = _sum_records_until(
                workbook,
                CUSTOMER_ABNORMAL_DEDUCTION_SHEET,
                date_column=0,
                amount_column=1,
                end_date=upload_date,
            )
            current_balance = total_paid - cumulative_fee - total_abnormal_deducted

        current_balance = round(current_balance, 2)
        return BalanceUploadRecord(
            customer=customer,
            today_fee=round(today_fee, 2),
            today_balance=current_balance,
            balance_date=latest_date,
            status="欠款" if current_balance < 0 else "充足",
            history_file=history_file,
        )
    finally:
        workbook.close()


def build_balance_upload_payload(
    preview: BalanceUploadPreview,
    *,
    app_version: str,
    uploaded_at: datetime | None = None,
) -> dict[str, Any]:
    timestamp = uploaded_at or datetime.now()
    return {
        "schemaVersion": "1.0",
        "appVersion": app_version,
        "uploadDate": preview.upload_date.isoformat(),
        "uploadedAt": timestamp.isoformat(timespec="seconds"),
        "sourceDir": str(preview.source_dir),
        "summary": {
            "customerCount": preview.customer_count,
            "totalTodayFee": preview.total_today_fee,
            "totalTodayBalance": preview.total_today_balance,
            "debtorCount": preview.debtor_count,
        },
        "records": [
            {
                "customer": record.customer,
                "todayFee": record.today_fee,
                "todayBalance": record.today_balance,
                "balanceDate": record.balance_date.isoformat(),
                "status": record.status,
            }
            for record in preview.records
        ],
    }


def upload_balance_payload(
    url: str,
    token: str,
    payload: dict[str, Any],
    *,
    timeout: int = 10,
    opener: Callable[..., Any] | None = None,
) -> BalanceUploadResult:
    url = url.strip()
    token = token.strip()
    if not url:
        return BalanceUploadResult(ok=False, status_code=None, message="上传接口地址未配置")
    if not token:
        return BalanceUploadResult(ok=False, status_code=None, message="上传密钥未配置")

    request = urllib.request.Request(
        url,
        data=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
        method="POST",
        headers={
            "Authorization": f"Bearer {token}",
            "Content-Type": "application/json",
        },
    )
    open_url = opener or urllib.request.urlopen
    try:
        with open_url(request, timeout=timeout) as response:
            status_code = int(
                response.status if hasattr(response, "status") else response.getcode()
            )
            body = response.read().decode("utf-8", errors="replace").strip()
    except urllib.error.HTTPError as exc:
        body = exc.read().decode("utf-8", errors="replace").strip()
        message = f"服务器返回错误：{exc.code}"
        if body:
            message = f"{message}，{body[:120]}"
        return BalanceUploadResult(ok=False, status_code=exc.code, message=message)
    except (OSError, urllib.error.URLError) as exc:
        return BalanceUploadResult(ok=False, status_code=None, message=f"上传失败：{exc}")

    if 200 <= status_code < 300:
        return BalanceUploadResult(ok=True, status_code=status_code, message="上传成功")
    message = f"服务器返回错误：{status_code}"
    if body:
        message = f"{message}，{body[:120]}"
    return BalanceUploadResult(ok=False, status_code=status_code, message=message)


def _sum_records_until(
    workbook: openpyxl.Workbook,
    sheet_name: str,
    *,
    date_column: int,
    amount_column: int,
    end_date: date,
) -> float:
    if sheet_name not in workbook.sheetnames:
        return 0.0
    sheet = workbook[sheet_name]
    total = 0.0
    for row in sheet.iter_rows(min_row=2, values_only=True):
        row_date = _parse_date(row[date_column] if len(row) > date_column else None)
        amount = _parse_number(row[amount_column] if len(row) > amount_column else None)
        if row_date is not None and row_date <= end_date and amount is not None:
            total += amount
    return total
