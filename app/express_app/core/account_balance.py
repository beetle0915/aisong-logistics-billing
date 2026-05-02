"""Read customer history workbooks for the account balance dashboard."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from express_app.core.calculator import (
    CUSTOMER_HISTORY_SHEET,
    CUSTOMER_PAYMENT_SHEET,
    customer_history_summary_file_name,
)


@dataclass(frozen=True)
class CustomerBalanceRecord:
    customer: str
    total_consumed: float
    total_paid: float
    current_balance: float
    last_date: date | None
    history_file: Path
    customer_dir: Path

    @property
    def status(self) -> str:
        return "欠款" if self.current_balance < 0 else "充足"


@dataclass(frozen=True)
class AccountBalanceDashboard:
    records: list[CustomerBalanceRecord] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def customer_count(self) -> int:
        return len(self.records)

    @property
    def total_consumed(self) -> float:
        return sum(record.total_consumed for record in self.records)

    @property
    def total_paid(self) -> float:
        return sum(record.total_paid for record in self.records)

    @property
    def total_balance(self) -> float:
        return sum(record.current_balance for record in self.records)

    @property
    def debtor_count(self) -> int:
        return sum(1 for record in self.records if record.current_balance < 0)


def collect_account_balance_dashboard(split_dir: Path) -> AccountBalanceDashboard:
    if not split_dir.exists():
        return AccountBalanceDashboard(errors=[f"客户明细目录不存在：{split_dir}"])

    records: list[CustomerBalanceRecord] = []
    errors: list[str] = []
    for customer_dir in sorted(path for path in split_dir.iterdir() if path.is_dir()):
        history_file = customer_dir / customer_history_summary_file_name(customer_dir.name)
        if not history_file.exists():
            continue
        try:
            records.append(read_customer_balance_record(customer_dir, history_file))
        except ValueError as exc:
            errors.append(f"{customer_dir.name}：{exc}")

    records.sort(
        key=lambda item: (
            item.last_date is None,
            -(item.last_date.toordinal() if item.last_date else 0),
            item.customer,
        )
    )
    return AccountBalanceDashboard(records=records, errors=errors)


def read_customer_balance_record(customer_dir: Path, history_file: Path) -> CustomerBalanceRecord:
    workbook = openpyxl.load_workbook(history_file, data_only=True, read_only=True)
    if CUSTOMER_HISTORY_SHEET not in workbook.sheetnames:
        raise ValueError(f"缺少 sheet：{CUSTOMER_HISTORY_SHEET}")

    history_sheet = workbook[CUSTOMER_HISTORY_SHEET]
    headers = _read_headers(history_sheet)
    required_headers = ["日期", "累计快递费用"]
    missing_headers = [header for header in required_headers if header not in headers]
    if missing_headers:
        raise ValueError(f"历史汇总缺少列：{'、'.join(missing_headers)}")

    last_date: date | None = None
    total_consumed = 0.0
    date_index = headers["日期"]
    total_fee_index = headers["累计快递费用"]
    for row in history_sheet.iter_rows(min_row=2, values_only=True):
        current_date = _parse_date(row[date_index])
        current_total = _parse_number(row[total_fee_index])
        if current_date is None and current_total is None:
            continue
        if current_date is not None:
            last_date = current_date
        if current_total is not None:
            total_consumed = current_total

    total_paid = _sum_payment_records(workbook)
    return CustomerBalanceRecord(
        customer=customer_dir.name,
        total_consumed=round(total_consumed, 2),
        total_paid=round(total_paid, 2),
        current_balance=round(total_paid - total_consumed, 2),
        last_date=last_date,
        history_file=history_file,
        customer_dir=customer_dir,
    )


def _read_headers(ws: openpyxl.worksheet.worksheet.Worksheet) -> dict[str, int]:
    values = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    return {str(value).strip(): index for index, value in enumerate(values) if value}


def _sum_payment_records(workbook: openpyxl.Workbook) -> float:
    if CUSTOMER_PAYMENT_SHEET not in workbook.sheetnames:
        return 0.0
    payment_sheet = workbook[CUSTOMER_PAYMENT_SHEET]
    total = 0.0
    for row in payment_sheet.iter_rows(min_row=2, values_only=True):
        amount = _parse_number(row[3] if len(row) > 3 else None)
        if amount is not None:
            total += amount
    return total


def _parse_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def _parse_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    return None
